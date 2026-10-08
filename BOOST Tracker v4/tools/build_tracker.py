"""Builds the BOOST Ontario tracker workbook (v4).

    python3 build_tracker.py            -> BOOST Tracker v4.xlsx (empty, ready for real use)
    python3 build_tracker.py --demo     -> BOOST Tracker v4 - DEMO.xlsx (fictional sample data)
"""
import datetime as dt
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

DEMO = "--demo" in sys.argv
OUT_DIR = Path(__file__).resolve().parent.parent
OUT = OUT_DIR / ("BOOST Tracker v4 - DEMO.xlsx" if DEMO else "BOOST Tracker v4.xlsx")
TODAY = dt.date(2026, 10, 9)  # demo data is written relative to the meeting date

# ---------------------------------------------------------------- sizes
P_FIRST, P_LAST = 5, 604          # People rows (600)
E_FIRST, E_LAST = 5, 104          # Events rows (100)
I_FIRST, I_LAST = 5, 1004         # New Sign-Ups rows (1000 form responses)
R_FIRST, R_LAST = 18, 97          # Event Report table rows (80)
F_FIRST, F_LAST = 7, 156          # Follow-Ups rows (150)
FORM_RANGE = "Sheet1!$A$1:$AZ$5000"

# ---------------------------------------------------------------- styles
NAVY, TEAL, GREY_HDR = "1F3864", "1B6E7A", "6B7785"
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
LIGHT = PatternFill("solid", fgColor="F3F5F8")
BANNER = PatternFill("solid", fgColor="E8EEF7")
DEMO_FILL = PatternFill("solid", fgColor="FCE4D6")
WHITE_B = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TITLE = Font(name="Calibri", size=16, bold=True, color=NAVY)
SUB = Font(name="Calibri", size=10, italic=True, color="595959")
BOLD = Font(name="Calibri", size=10, bold=True)
BODY = Font(name="Calibri", size=10)
H2 = Font(name="Calibri", size=12, bold=True, color=NAVY)
THIN = Side(style="thin", color="D0D5DD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
UNLOCK = Protection(locked=False)
DATE, DATETIME = "yyyy-mm-dd", "yyyy-mm-dd hh:mm"


def fill(c):
    return PatternFill("solid", fgColor=c)


def header(ws, row, cols, color=NAVY, height=32):
    """cols: list of (letter, text, width, color-or-None)."""
    for letter, text, width, col_color in cols:
        c = ws[f"{letter}{row}"]
        c.value = text
        c.font = WHITE_B
        c.fill = fill(col_color or color)
        c.alignment = CENTER
        c.border = BOX
        if width:
            ws.column_dimensions[letter].width = width
    ws.row_dimensions[row].height = height


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws.row_dimensions[1].height = 26
    if sub:
        ws["A2"] = sub
        ws["A2"].font = SUB
    if DEMO:
        ws["A2"].value = "SAMPLE DATA — every person in this file is fictional.  " + (sub or "")
        ws["A2"].font = Font(name="Calibri", size=10, bold=True, color="C55A11")
    ws.sheet_view.showGridLines = False


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def protect(ws):
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def dv_list(ws, rng, source, allow_other=False):
    d = DataValidation(type="list", formula1=source, allow_blank=True)
    d.showErrorMessage = not allow_other
    if not allow_other:
        d.errorTitle = "Pick from the list"
        d.error = "Use the dropdown so the counts stay consistent. New values go at the bottom of the Lists tab."
    ws.add_data_validation(d)
    d.add(rng)


wb = Workbook()
wb.remove(wb.active)
order = ["Start Here", "People", "New Sign-Ups", "Follow-Ups", "Events", "Event Report",
         "Dashboard", "Contact Card", "Setup", "Lists", "Sheet1"]
S = {n: wb.create_sheet(n) for n in order}

# =================================================================== Lists
ws = S["Lists"]
title(ws, "Lists", "Dropdown vocabulary. Add new values at the bottom of a column — the dropdowns grow with it. Don't rename existing values once in use.")
LISTS = {
    "L_Type": ("Contact Type", ["Trade", "Employer & Union", "College & Agency", "Other"]),
    "L_Region": ("Region", ["Northeastern Ontario", "Northwestern Ontario", "Eastern Ontario", "Central Ontario",
                            "Southwestern Ontario", "Greater Toronto Area", "Other / not sure"]),
    "L_Situation": ("Situation", ["Jobseeker", "Incumbent worker", "Journeyperson", "Student", "Prefer not to say"]),
    "L_YNU": ("Registered Apprentice", ["Yes", "No", "Not sure"]),
    "L_Service": ("Service Level", ["Full-Service", "Light-Touch", "Not decided"]),
    "L_Need": ("Main Need", ["Finding a sponsor / employer", "Exam preparation", "Tutoring / upgrading",
                             "Reading, writing or math", "Learning disability testing", "Where to start",
                             "Info about BOOST", "Hiring an apprentice", "Partnership / referrals", "Other"]),
    "L_Stage": ("Stage", ["New", "Contacted", "Intake booked", "Active", "Referred out", "No response", "Closed"]),
    "L_Tick": ("Tick", ["✓"]),
    "L_YesNo": ("Consent", ["Yes", "No"]),
    "L_EventType": ("Event Type", ["SOY Boot Camp", "Level Up", "Career fair", "Information session",
                                   "Presentation", "Community event", "Other"]),
}
for i, (nm, (label, values)) in enumerate(LISTS.items()):
    col = chr(ord("A") + i)
    header(ws, 4, [(col, label, 24, None)])
    for j, v in enumerate(values):
        ws[f"{col}{5 + j}"] = v
    name(wb, nm, f"OFFSET(Lists!${col}$5,0,0,MAX(1,COUNTA(Lists!${col}$5:${col}$60)),1)")

# =================================================================== Sheet1 (form placeholder)
FORM_HEADERS = ["Id", "Start time", "Completion time", "Email", "Name", "Last modified time",
                "Who are you?", "Full name", "Email address", "Phone number", "City or town",
                "Company name", "Organization name",
                "What region of Ontario are you in?", "Which best describes you right now?",
                "Are you currently a registered apprentice?", "What trade are you in, or interested in?",
                "Do you identify with any of the following? (optional)",
                "What would be most helpful to you right now?", "Consent"]
ws = S["Sheet1"]
for i, h in enumerate(FORM_HEADERS, 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = BOLD
    c.fill = fill("DDEBF7")
    ws.column_dimensions[c.column_letter].width = 18
ws.freeze_panes = "A2"

# =================================================================== Setup
ws = S["Setup"]
title(ws, "Setup", "Tells the tracker where to find each answer in the form's response sheet (Sheet1). Only the yellow cells are yours to change.")
ws["A3"] = ("Each yellow cell holds a piece of the form's column heading. The tracker looks for an exact match first, "
            "then for any heading that contains it. Check column E shows the heading you meant. Blank = not used.")
ws["A3"].font = SUB
ws.merge_cells("A3:F3")
ws.row_dimensions[3].height = 30
ws["A3"].alignment = WRAP
header(ws, 5, [("A", "Key", 6, GREY_HDR), ("B", "What it is", 30, None), ("C", "Heading text to look for", 34, None),
               ("D", "Column #", 10, GREY_HDR), ("E", "Heading found", 46, GREY_HDR), ("F", "Used for", 44, GREY_HDR)])
MAP = [
    ("Id", "Response ID", "Id", "The number you type into People → Form Response ID"),
    ("Time", "Submitted (completion time)", "Completion time", "Registered On, and the event the response belongs to"),
    ("Name", "Full name", "Full name", "Initials only. The full name is shown on Contact Card, nowhere else"),
    ("Email", "Email address", "Email address", "Contact Card only"),
    ("Phone", "Phone number", "Phone", "Contact Card only"),
    ("City", "City or town", "City", "Contact Card only"),
    ("Who", "Who are you (branch question)", "Who are you", "Suggests Contact Type, using the keyword table below"),
    ("Org1", "Company name", "Company", "Organization (employers)"),
    ("Org2", "Organization name", "Organization name", "Organization (colleges, agencies, partners)"),
    ("Region", "Region", "region", "Region"),
    ("Sit", "Which best describes you", "describes you", "Situation, using the keyword table below"),
    ("Appr", "Registered apprentice?", "registered apprentice", "Registered Apprentice"),
    ("Trade", "Trade", "What trade", "Trade"),
    ("SelfID", "Self-identification (optional)", "identify", "Dashboard totals ONLY — never copied onto a person's row"),
    ("Need", "Most helpful right now", "most helpful", "Shown as a hint on New Sign-Ups"),
    ("Consent", "Consent", "Consent", "Consent (any answer containing “yes” counts as Yes)"),
    ("Event", "Event (pre-filled link, if you add one)", "signing up today", "Event ID, when the QR carries the event. Otherwise the date decides"),
]
for i, (key, what, default, used) in enumerate(MAP):
    r = 6 + i
    ws[f"A{r}"] = key
    ws[f"B{r}"] = what
    ws[f"C{r}"] = default
    ws[f"C{r}"].fill = INPUT_FILL
    ws[f"C{r}"].protection = UNLOCK
    ws[f"D{r}"] = f'=IF(C{r}="","",IFERROR(MATCH(C{r},FormHdr,0),IFERROR(MATCH("*"&C{r}&"*",FormHdr,0),"")))'
    ws[f"E{r}"] = f'=IF(C{r}="","— not used",IF(D{r}="","✗ not found in Sheet1 — check the wording",INDEX(FormHdr,D{r})))'
    ws[f"F{r}"] = used
    for col in "ABCDEF":
        ws[f"{col}{r}"].border = BOX
        ws[f"{col}{r}"].font = BODY
        ws[f"{col}{r}"].alignment = WRAP
    name(wb, f"f_{key}", f"Setup!$D${r}")
name(wb, "FormData", FORM_RANGE)
name(wb, "FormHdr", "Sheet1!$A$1:$AZ$1")
ws.conditional_formatting.add("E6:E22", FormulaRule(formula=['LEFT($E6,1)="✗"'], font=Font(color="C00000", bold=True)))

r0 = 25
ws[f"A{r0}"] = "Settings"
ws[f"A{r0}"].font = H2
ws[f"B{r0+1}"] = "“Due soon” means within this many days"
ws[f"C{r0+1}"] = 7
ws[f"C{r0+1}"].fill = INPUT_FILL
ws[f"C{r0+1}"].protection = UNLOCK
name(wb, "DueSoon", f"Setup!$C${r0+1}")

r1 = 29
ws[f"A{r1}"] = "Commitments shown on the Dashboard"
ws[f"A{r1}"].font = H2
ws[f"A{r1+1}"] = ("Targets copied from the BOOST Performance Commitments 2026-27 (V2.0) as quoted in the sign-up form spec. "
                  "Check them against the source before relying on them; change any yellow cell.")
ws[f"A{r1+1}"].font = SUB
header(ws, r1 + 2, [("B", "Commitment", None, None), ("C", "Target", None, None)])
TARGETS = [("KPI 1 — Total connections", None), ("KPI 2 — Underrepresented groups", 106),
           ("KPI 3 — Jobseekers", 118), ("KPI 4 — Incumbent workers", 102),
           ("KPI 6 — Registered apprentices", 60), ("KPI 26 — Persons with a disability", 69),
           ("Full-Service participants", 280), ("Light-Touch participants", 100)]
for i, (lbl, tgt) in enumerate(TARGETS):
    r = r1 + 3 + i
    ws[f"B{r}"] = lbl
    ws[f"C{r}"] = tgt
    ws[f"C{r}"].fill = INPUT_FILL
    ws[f"C{r}"].protection = UNLOCK
    ws[f"B{r}"].border = ws[f"C{r}"].border = BOX
TARGET_ROW = {lbl: r1 + 3 + i for i, (lbl, _) in enumerate(TARGETS)}

r2 = 43
ws[f"A{r2}"] = "Keyword tables — how form answers become dropdown values"
ws[f"A{r2}"].font = H2
ws[f"A{r2+1}"] = ("If an answer contains the keyword, it gets the value beside it. When several match, the LOWEST row wins — "
                  "so put more specific keywords further down. Add rows in the blank yellow cells.")
ws[f"A{r2+1}"].font = SUB
header(ws, r2 + 2, [("B", "Who are you → keyword", None, None), ("C", "Contact Type", None, None),
                    ("E", "Situation → keyword", None, None), ("F", "Situation", None, None)])
WHO = [("trade", "Trade"), ("apprentic", "Trade"), ("student", "Trade"), ("job", "Trade"),
       ("employ", "Employer & Union"), ("union", "Employer & Union"), ("sponsor", "Employer & Union"),
       ("business", "Employer & Union"), ("college", "College & Agency"), ("school", "College & Agency"),
       ("agency", "College & Agency"), ("organization", "College & Agency"), ("partner", "College & Agency"),
       ("government", "College & Agency"), ("parent", "Other"), ("other", "Other")]
SIT = [("looking for work", "Jobseeker"), ("not working", "Jobseeker"), ("journeyperson", "Journeyperson"),
       ("not a journeyperson", "Incumbent worker"), ("student", "Student"), ("prefer not", "Prefer not to say")]
KW_FIRST, KW_LAST = r2 + 3, r2 + 3 + 21
for i in range(KW_LAST - KW_FIRST + 1):
    r = KW_FIRST + i
    for col in "BCEF":
        ws[f"{col}{r}"].fill = INPUT_FILL
        ws[f"{col}{r}"].protection = UNLOCK
        ws[f"{col}{r}"].border = BOX
    if i < len(WHO):
        ws[f"B{r}"], ws[f"C{r}"] = WHO[i]
    if i < len(SIT):
        ws[f"E{r}"], ws[f"F{r}"] = SIT[i]
name(wb, "WhoKW", f"Setup!$B${KW_FIRST}:$B${KW_LAST}")
name(wb, "WhoVal", f"Setup!$C${KW_FIRST}:$C${KW_LAST}")
name(wb, "SitKW", f"Setup!$E${KW_FIRST}:$E${KW_LAST}")
name(wb, "SitVal", f"Setup!$F${KW_FIRST}:$F${KW_LAST}")
ws.column_dimensions["E"].width = 46
protect(ws)

# =================================================================== Events
ws = S["Events"]
title(ws, "Events", "One row per event. Give it an ID and a DATE before the day — sign-ups find their event by the date.")
ws["A3"] = "Yellow = yours to fill.  Grey = counts itself."
ws["A3"].font = SUB
header(ws, 4, [("A", "Event ID", 10, None), ("B", "Date", 12, None), ("C", "Event Name", 30, None),
               ("D", "Event Type", 18, None), ("E", "Host / Partner", 24, None), ("F", "City", 14, None),
               ("G", "Region", 20, None), ("H", "Staff", 14, None),
               ("I", "Form Sign-Ups", 10, GREY_HDR), ("J", "People Logged", 10, GREY_HDR),
               ("K", "Sign-Ups Not Yet Logged", 12, GREY_HDR), ("L", "Follow-Ups Open", 10, GREY_HDR),
               ("M", "Report Sent", 11, None), ("N", "What Worked / Notes", 50, None)])
PB, PE, PS, PP, PR, PO, PT, PZ = (f"People!${c}${P_FIRST}:${c}${P_LAST}" for c in "BESPROTZ")
IE, IG = f"'New Sign-Ups'!$E${I_FIRST}:$E${I_LAST}", f"'New Sign-Ups'!$G${I_FIRST}:$G${I_LAST}"
for r in range(E_FIRST, E_LAST + 1):
    ws[f"I{r}"] = f'=IF($A{r}="","",COUNTIF({IE},$A{r}))'
    ws[f"J{r}"] = f'=IF($A{r}="","",COUNTIFS({PB},$A{r},{PZ},1))'
    ws[f"K{r}"] = f'=IF($A{r}="","",COUNTIFS({IE},$A{r},{IG},"→ add"))'
    ws[f"L{r}"] = f'=IF($A{r}="","",COUNTIFS({PB},$A{r},{PP},"✓",{PR},""))'
    ws[f"B{r}"].number_format = DATE
    ws[f"M{r}"].number_format = DATE
    for col in "ABCDEFGHIJKLMN":
        c = ws[f"{col}{r}"]
        c.border = BOX
        c.font = BODY
        c.alignment = WRAP if col in "CEN" else Alignment(vertical="top", horizontal="center" if col in "ABIJKLM" else None)
        if col in "IJKL":
            c.fill = LIGHT
ws.freeze_panes = "C5"
ws.auto_filter.ref = f"A4:N{E_LAST}"
dv_list(ws, f"D{E_FIRST}:D{E_LAST}", "L_EventType", allow_other=True)
dv_list(ws, f"G{E_FIRST}:G{E_LAST}", "L_Region")
rng = f"A{E_FIRST}:N{E_LAST}"
ws.conditional_formatting.add(f"B{E_FIRST}:B{E_LAST}", FormulaRule(formula=[f'AND($A{E_FIRST}<>"",$B{E_FIRST}="")'], fill=fill("F8CBAD")))
ws.conditional_formatting.add(f"B{E_FIRST}:B{E_LAST}", FormulaRule(formula=[f'AND($B{E_FIRST}<>"",COUNTIF($B${E_FIRST}:$B${E_LAST},$B{E_FIRST})>1)'], fill=fill("FFD966")))
ws.conditional_formatting.add(f"A{E_FIRST}:A{E_LAST}", FormulaRule(formula=[f'AND($A{E_FIRST}<>"",COUNTIF($A${E_FIRST}:$A${E_LAST},$A{E_FIRST})>1)'], fill=fill("F8CBAD")))
ws.conditional_formatting.add(f"K{E_FIRST}:K{E_LAST}", FormulaRule(formula=[f'N($K{E_FIRST})>0'], font=Font(color="C55A11", bold=True)))
name(wb, "EvID", f"Events!$A${E_FIRST}:$A${E_LAST}")
name(wb, "EvDate", f"Events!$B${E_FIRST}:$B${E_LAST}")
name(wb, "EvName", f"Events!$C${E_FIRST}:$C${E_LAST}")

# =================================================================== People
ws = S["People"]
title(ws, "People", "One row per person you speak with. Initials and a number only — names, emails and phones stay in the form.")
ws["A3"] = "◀ ID"
ws["B3"] = "AUTO — fills from the form once you type the Form Response ID (col U).  Type over any of these to change that row."
ws["K3"] = "YOURS — the conversation"
ws["U3"] = "LINK"
for cell in ("A3", "B3", "K3", "U3"):
    ws[cell].font = Font(name="Calibri", size=9, bold=True, color=GREY_HDR)
AUTO, MINE = TEAL, NAVY
header(ws, 4, [("A", "Participant #", 10, GREY_HDR), ("B", "Event ID", 9, AUTO), ("C", "Registered On", 15, AUTO),
               ("D", "Initials", 8, AUTO), ("E", "Contact Type", 16, AUTO), ("F", "Organization", 22, AUTO),
               ("G", "Region", 18, AUTO), ("H", "Situation", 15, AUTO), ("I", "Reg. Apprentice", 10, AUTO),
               ("J", "Trade", 16, AUTO), ("K", "Service Level", 12, MINE), ("L", "Main Need", 22, MINE),
               ("M", "Stage", 13, MINE), ("N", "Conversation Notes", 60, MINE), ("O", "Materials Given", 9, MINE),
               ("P", "Follow-Up Needed", 9, MINE), ("Q", "Follow-Up Due", 11, MINE), ("R", "Follow-Up Done", 11, MINE),
               ("S", "Follow-Up Status", 13, GREY_HDR), ("T", "Consent", 9, AUTO), ("U", "Form Response ID", 10, "C55A11"),
               ("V", "Form Check", 10, GREY_HDR),
               ("W", "_form row", 8, GREY_HDR), ("X", "_report", 8, GREY_HDR), ("Y", "_follow-up", 8, GREY_HDR),
               ("Z", "_in use", 8, GREY_HDR)], height=40)
FD = lambda r, key: f"INDEX(FormData,$W{r},f_{key})"  # noqa: E731


def initials(expr):
    n = f"TRIM({expr})"
    return (f'UPPER(LEFT({n},1)&IFERROR(MID({n},FIND("^",SUBSTITUTE({n}," ","^",'
            f'LEN({n})-LEN(SUBSTITUTE({n}," ",""))))+1,1),""))')


def keyword(ans, kw, val):
    k = (f"SUMPRODUCT(MAX(({kw}<>\"\")*ISNUMBER(SEARCH({kw},{ans}))"
         f"*(ROW({kw})-MIN(ROW({kw}))+1)))")
    return f'IF({k}=0,"",INDEX({val},{k}))'


def event_lookup(rowref, timecell, fallback):
    ev = f"INDEX(FormData,{rowref},f_Event)"
    return (f'IFERROR(INDEX(EvID,MATCH({ev},EvID,0)),IFERROR(INDEX(EvID,MATCH({ev},EvName,0)),'
            f'IFERROR(INDEX(EvID,MATCH(INT({timecell}),EvDate,0)),{fallback})))')


for r in range(P_FIRST, P_LAST + 1):
    ws[f"A{r}"] = f"P-{r - P_FIRST + 1:04d}"
    ws[f"B{r}"] = f'=IF(AND($U{r}="",$C{r}=""),"",{event_lookup(f"$W{r}", f"$C{r}", chr(34)*2)})'
    ws[f"C{r}"] = f'=IF($W{r}="","",{FD(r, "Time")})'
    ws[f"D{r}"] = f'=IF($W{r}="","",IFERROR({initials(FD(r, "Name"))},""))'
    ws[f"E{r}"] = f'=IF($W{r}="","",IFERROR({keyword(FD(r, "Who") + chr(38) + chr(34)*2, "WhoKW", "WhoVal")},""))'
    ws[f"F{r}"] = (f'=IF($W{r}="","",TRIM(IFERROR({FD(r, "Org1")}&"","")&" "&IFERROR({FD(r, "Org2")}&"","")))')
    ws[f"G{r}"] = f'=IF($W{r}="","",IFERROR({FD(r, "Region")}&"",""))'
    ws[f"H{r}"] = f'=IF($W{r}="","",IFERROR({keyword(FD(r, "Sit") + chr(38) + chr(34)*2, "SitKW", "SitVal")},""))'
    ws[f"I{r}"] = f'=IF($W{r}="","",IFERROR({FD(r, "Appr")}&"",""))'
    ws[f"J{r}"] = f'=IF($W{r}="","",IFERROR({FD(r, "Trade")}&"",""))'
    ws[f"S{r}"] = (f'=IF($P{r}<>"✓","",IF($R{r}<>"","Done",IF($Q{r}="","Open — no date",'
                   f'IF($Q{r}<TODAY(),"Overdue",IF($Q{r}-TODAY()<=DueSoon,"Due soon","Scheduled")))))')
    ws[f"T{r}"] = (f'=IF($W{r}="","",IFERROR(IF({FD(r, "Consent")}&""="","",'
                   f'IF(ISNUMBER(SEARCH("yes",{FD(r, "Consent")})),"Yes","No")),""))')
    ws[f"V{r}"] = f'=IF($U{r}="","",IF($W{r}="","not found","✓"))'
    ws[f"W{r}"] = (f'=IF($U{r}="","",IFERROR(MATCH($U{r},INDEX(FormData,0,f_Id),0),'
                   f'IFERROR(MATCH(--$U{r},INDEX(FormData,0,f_Id),0),IFERROR(MATCH($U{r}&"",INDEX(FormData,0,f_Id),0),""))))')
    ws[f"X{r}"] = f'=IF(AND($B{r}<>"",$B{r}=\'Event Report\'!$C$3),ROW(),"")'
    ws[f"Y{r}"] = f'=IF(AND($P{r}="✓",$R{r}=""),IF($Q{r}="",90000,$Q{r})+ROW()/100000,"")'
    ws[f"Z{r}"] = f'=IF(OR($U{r}<>"",$D{r}<>"",$E{r}<>"",$N{r}<>""),1,"")'
    for col in "ABCDEFGHIJKLMNOPQRSTUV":
        c = ws[f"{col}{r}"]
        c.border = BOX
        c.font = BODY
        c.alignment = WRAP if col in "FLN" else Alignment(vertical="top", horizontal="center" if col in "ABDIOPQRSTUV" else None)
    ws[f"A{r}"].fill = LIGHT
    ws[f"S{r}"].fill = LIGHT
    ws[f"V{r}"].fill = LIGHT
    ws[f"C{r}"].number_format = DATETIME
    ws[f"Q{r}"].number_format = DATE
    ws[f"R{r}"].number_format = DATE
for col in "WXYZ":
    ws.column_dimensions[col].hidden = True
ws.freeze_panes = "E5"
ws.auto_filter.ref = f"A4:V{P_LAST}"
rows = lambda c: f"{c}{P_FIRST}:{c}{P_LAST}"  # noqa: E731
dv_list(ws, rows("B"), "EvID")
dv_list(ws, rows("E"), "L_Type")
dv_list(ws, rows("G"), "L_Region")
dv_list(ws, rows("H"), "L_Situation")
dv_list(ws, rows("I"), "L_YNU")
dv_list(ws, rows("K"), "L_Service")
dv_list(ws, rows("L"), "L_Need")
dv_list(ws, rows("M"), "L_Stage")
dv_list(ws, f"O{P_FIRST}:P{P_LAST}", "L_Tick")
dv_list(ws, rows("T"), "L_YesNo")
dd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2024,1,1)", allow_blank=True)
dd.error = "Type a date, e.g. 2026-10-15"
ws.add_data_validation(dd)
dd.add(f"Q{P_FIRST}:R{P_LAST}")
f = P_FIRST
cf = ws.conditional_formatting
cf.add(rows("B"), FormulaRule(formula=[f'AND($Z{f}=1,$B{f}="")'], fill=fill("F8CBAD")))
cf.add(rows("B"), FormulaRule(formula=[f'AND(ISNUMBER($C{f}),COUNTIF(EvDate,INT(N($C{f})))>1)'], fill=fill("FFD966")))
cf.add(rows("S"), FormulaRule(formula=[f'$S{f}="Overdue"'], fill=fill("F8CBAD"), font=Font(color="9C0006", bold=True)))
cf.add(rows("S"), FormulaRule(formula=[f'$S{f}="Due soon"'], fill=fill("FFEB9C"), font=Font(color="7F6000", bold=True)))
cf.add(rows("S"), FormulaRule(formula=[f'$S{f}="Done"'], font=Font(color="548235")))
cf.add(rows("V"), FormulaRule(formula=[f'$V{f}="not found"'], font=Font(color="C00000", bold=True), fill=fill("F8CBAD")))
cf.add(rows("T"), FormulaRule(formula=[f'$T{f}="No"'], font=Font(color="C00000", bold=True)))
cf.add(rows("F"), FormulaRule(formula=[f'OR($E{f}="Trade",$E{f}="Other")'], fill=fill("EDEDED")))
cf.add(f"H{P_FIRST}:K{P_LAST}", FormulaRule(formula=[f'OR($E{f}="Employer & Union",$E{f}="College & Agency")'], fill=fill("EDEDED")))
cf.add(f"A{P_FIRST}:A{P_LAST}", FormulaRule(formula=[f'$Z{f}=1'], font=Font(bold=True, color=NAVY)))

# =================================================================== New Sign-Ups
ws = S["New Sign-Ups"]
title(ws, "New Sign-Ups", "Everyone who has filled in the form, newest at the bottom. Anyone marked → add isn't on People yet.")
ws["A3"] = "Not yet on People:"
ws["A3"].font = BOLD
ws["C3"] = f'=COUNTIF(G{I_FIRST}:G{I_LAST},"→ add")'
ws["C3"].font = Font(name="Calibri", size=14, bold=True, color="C55A11")
ws["D3"] = "→ Filter column G to “→ add”, then type each Response ID into a blank row on People (column U)."
ws["D3"].font = SUB
header(ws, 4, [("A", "#", 6, GREY_HDR), ("B", "Response ID", 11, None), ("C", "Submitted", 16, None),
               ("D", "Initials", 8, None), ("E", "Event (by date)", 16, None), ("F", "Contact Type (from form)", 18, None),
               ("G", "On People?", 12, None), ("H", "_row", 6, GREY_HDR), ("I", "Wants help with (from form)", 46, None)])
for r in range(I_FIRST, I_LAST + 1):
    n = r - I_FIRST + 1
    ws[f"A{r}"] = n
    ws[f"H{r}"] = f'=IFERROR(IF(INDEX(FormData,{n + 1},f_Id)&""="","",{n + 1}),"")'
    ws[f"B{r}"] = f'=IF($H{r}="","",INDEX(FormData,$H{r},f_Id))'
    ws[f"C{r}"] = f'=IF($H{r}="","",IFERROR(INDEX(FormData,$H{r},f_Time),""))'
    ws[f"D{r}"] = f'=IF($H{r}="","",IFERROR({initials(f"INDEX(FormData,$H{r},f_Name)")},""))'
    ws[f"E{r}"] = f'=IF($H{r}="","",{event_lookup(f"$H{r}", f"$C{r}", chr(34) + "(no event that day)" + chr(34))})'
    ws[f"F{r}"] = (f'=IF($H{r}="","",IFERROR({keyword(f"INDEX(FormData,$H{r},f_Who)&{chr(34)*2}", "WhoKW", "WhoVal")},""))')
    ws[f"G{r}"] = f'=IF($H{r}="","",IF(COUNTIF(People!$U${P_FIRST}:$U${P_LAST},$B{r})+COUNTIF(People!$U${P_FIRST}:$U${P_LAST},$B{r}&"")>0,"✓ logged","→ add"))'
    ws[f"I{r}"] = f'=IF($H{r}="","",IFERROR(INDEX(FormData,$H{r},f_Need)&"",""))'
    ws[f"C{r}"].number_format = DATETIME
    for col in "ABCDEFGI":
        ws[f"{col}{r}"].font = BODY
        ws[f"{col}{r}"].border = BOX
        ws[f"{col}{r}"].alignment = Alignment(vertical="top", horizontal="center" if col in "ABDG" else None)
ws.column_dimensions["H"].hidden = True
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:I{I_LAST}"
ws.conditional_formatting.add(f"G{I_FIRST}:G{I_LAST}", FormulaRule(formula=[f'$G{I_FIRST}="→ add"'], fill=fill("FCE4D6"), font=Font(color="C55A11", bold=True)))
ws.conditional_formatting.add(f"G{I_FIRST}:G{I_LAST}", FormulaRule(formula=[f'$G{I_FIRST}="✓ logged"'], font=Font(color="548235")))
ws.conditional_formatting.add(f"E{I_FIRST}:E{I_LAST}", FormulaRule(formula=[f'LEFT($E{I_FIRST},1)="("'], font=Font(color="9C0006", italic=True)))
protect(ws)

# =================================================================== Follow-Ups
ws = S["Follow-Ups"]
title(ws, "Follow-Ups", "Everyone still owed a call, soonest first. Fills itself from People — tick Follow-Up Needed there, add a due date, and enter the date when it's done.")
ws["A3"] = "Overdue"
ws["B3"] = f'=COUNTIF(People!$S${P_FIRST}:$S${P_LAST},"Overdue")'
ws["C3"] = "Due soon"
ws["D3"] = f'=COUNTIF(People!$S${P_FIRST}:$S${P_LAST},"Due soon")'
ws["E3"] = "All open"
ws["F3"] = f'=COUNT(People!$Y${P_FIRST}:$Y${P_LAST})'
for c, color in (("B3", "C00000"), ("D3", "BF8F00"), ("F3", NAVY)):
    ws[c].font = Font(name="Calibri", size=14, bold=True, color=color)
for c in ("A3", "C3", "E3"):
    ws[c].font = BOLD
    ws[c].alignment = Alignment(horizontal="right")
ws["A4"] = "Edit on People, not here. Find the person with their Participant #."
ws["A4"].font = SUB
header(ws, 6, [("A", "Due", 11, None), ("B", "Status", 12, None), ("C", "Participant #", 11, None), ("D", "Initials", 8, None),
               ("E", "Event", 9, None), ("F", "Contact Type", 16, None), ("G", "Main Need", 22, None), ("H", "Stage", 13, None),
               ("I", "Latest notes", 70, None), ("J", "_key", 8, GREY_HDR), ("K", "_row", 6, GREY_HDR)])
for r in range(F_FIRST, F_LAST + 1):
    k = r - F_FIRST + 1
    ws[f"J{r}"] = f'=IFERROR(SMALL(People!$Y${P_FIRST}:$Y${P_LAST},{k}),"")'
    ws[f"K{r}"] = f'=IF($J{r}="","",ROUND(MOD($J{r},1)*100000,0))'
    g = lambda col: f"INDEX(People!${col}:${col},$K{r})"  # noqa: E731
    ws[f"A{r}"] = f'=IF($K{r}="","",IF({g("Q")}="","no date",{g("Q")}))'
    ws[f"B{r}"] = f'=IF($K{r}="","",{g("S")})'
    for col, src in zip("CDEFGH", "ADBELM"):
        ws[f"{col}{r}"] = f'=IF($K{r}="","",{g(src)}&"")'
    ws[f"I{r}"] = f'=IF($K{r}="","",{g("N")}&"")'
    ws[f"A{r}"].number_format = DATE
    for col in "ABCDEFGHI":
        ws[f"{col}{r}"].font = BODY
        ws[f"{col}{r}"].alignment = WRAP if col == "I" else Alignment(vertical="top", horizontal="center" if col in "ABCDE" else None)
ws.column_dimensions["J"].hidden = True
ws.column_dimensions["K"].hidden = True
ws.freeze_panes = "A7"
ws.conditional_formatting.add(f"A{F_FIRST}:B{F_LAST}", FormulaRule(formula=[f'$B{F_FIRST}="Overdue"'], fill=fill("F8CBAD"), font=Font(color="9C0006", bold=True)))
ws.conditional_formatting.add(f"A{F_FIRST}:B{F_LAST}", FormulaRule(formula=[f'$B{F_FIRST}="Due soon"'], fill=fill("FFEB9C"), font=Font(color="7F6000", bold=True)))
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "6:6"
ws.conditional_formatting.add(f"A{F_FIRST}:I{F_LAST}", FormulaRule(formula=[f'$K{F_FIRST}<>""'], border=BOX))
ws.defined_names["_xlnm.Print_Area"] = DefinedName(
    "_xlnm.Print_Area", attr_text=f"OFFSET('Follow-Ups'!$A$1,0,0,{F_FIRST - 1}+MAX(1,COUNT('Follow-Ups'!$K${F_FIRST}:$K${F_LAST})),9)")
protect(ws)

# =================================================================== Event Report
ws = S["Event Report"]
title(ws, "BOOST Ontario — Event Report", "Pick an Event ID in the yellow box. Everything else fills itself. Write your note, then File → Print → Save as PDF.")
ws["B3"] = "Event ID ▸"
ws["B3"].font = Font(name="Calibri", size=12, bold=True)
ws["B3"].alignment = Alignment(horizontal="right", vertical="center")
ws["C3"].fill = INPUT_FILL
ws["C3"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
ws["C3"].protection = UNLOCK
ws["C3"].border = Border(left=Side("medium", NAVY), right=Side("medium", NAVY), top=Side("medium", NAVY), bottom=Side("medium", NAVY))
ws["C3"].alignment = CENTER
ws.row_dimensions[3].height = 24
dv_list(ws, "C3", "EvID")
ev = lambda col: f'=IF($C$3="","",IFERROR(INDEX(Events!${col}${E_FIRST}:${col}${E_LAST},MATCH($C$3,EvID,0))&"",""))'  # noqa: E731
details = [("Event", ev("C")), ("Date", f'=IF($C$3="","",IFERROR(INDEX(EvDate,MATCH($C$3,EvID,0)),""))'),
           ("Type", ev("D")), ("Host / Partner", ev("E")), ("Location", f'=IF($C$3="","",TRIM({ev("F")[1:]}&"  ·  "&{ev("G")[1:]}))'),
           ("Staff", ev("H"))]
for i, (lbl, frm) in enumerate(details):
    r = 5 + i
    ws[f"B{r}"] = lbl
    ws[f"B{r}"].font = BOLD
    ws[f"C{r}"] = frm
    ws[f"C{r}"].font = BODY
    ws.merge_cells(f"C{r}:E{r}")
ws["C6"].number_format = "dddd, mmmm d, yyyy"
ws["C6"].alignment = Alignment(horizontal="left")
cnt = lambda extra: f'=IF($C$3="","",COUNTIFS({PB},$C$3,{PZ},1{extra}))'  # noqa: E731
stats = [("Form sign-ups", f'=IF($C$3="","",COUNTIF({IE},$C$3))'),
         ("People logged", cnt("")),
         ("  Trade", cnt(f',{PE},"Trade"')), ("  Employers & Unions", cnt(f',{PE},"Employer & Union"')),
         ("  Colleges & Agencies", cnt(f',{PE},"College & Agency"')), ("  Other", cnt(f',{PE},"Other"')),
         ("Materials given", cnt(f',{PO},"✓"')), ("Follow-ups open", cnt(f',{PP},"✓",{PR},""')),
         ("Consented to contact", cnt(f',{PT},"Yes"'))]
for i, (lbl, frm) in enumerate(stats):
    r = 5 + i
    ws[f"G{r}"] = lbl
    ws[f"G{r}"].font = BODY if lbl.startswith("  ") else BOLD
    ws[f"H{r}"] = frm
    ws[f"H{r}"].font = BOLD
    ws[f"H{r}"].alignment = Alignment(horizontal="center")
ws["B14"] = "Notes for the reader"
ws["B14"].font = H2
ws.merge_cells("B15:I15")
ws["B15"].alignment = WRAP
ws["B15"].fill = INPUT_FILL
ws["B15"].protection = UNLOCK
ws["B15"].border = BOX
ws.row_dimensions[15].height = 80
header(ws, 17, [("A", "#", 4, None), ("B", "Participant #", 12, None), ("C", "Initials", 8, None),
                ("D", "Contact Type", 16, None), ("E", "Organization", 20, None), ("F", "Main Need", 20, None),
                ("G", "Stage", 16, None), ("H", "Follow-Up", 18, None), ("I", "Conversation Notes", 62, None),
                ("K", "_row", 6, GREY_HDR)])
for r in range(R_FIRST, R_LAST + 1):
    k = r - R_FIRST + 1
    ws[f"K{r}"] = f'=IFERROR(SMALL(People!$X${P_FIRST}:$X${P_LAST},{k}),"")'
    g = lambda col: f"INDEX(People!${col}:${col},$K{r})"  # noqa: E731
    ws[f"A{r}"] = f'=IF($K{r}="","",{k})'
    for col, src in zip("BCDEFG", "ADEFLM"):
        ws[f"{col}{r}"] = f'=IF($K{r}="","",{g(src)}&"")'
    ws[f"H{r}"] = (f'=IF($K{r}="","",IF({g("P")}<>"✓","—",IF({g("R")}<>"","Done "&TEXT({g("R")},"yyyy-mm-dd"),'
                   f'"Yes"&IF({g("Q")}<>""," — by "&TEXT({g("Q")},"yyyy-mm-dd"),""))))')
    ws[f"I{r}"] = f'=IF($K{r}="","",{g("N")}&"")'
    for col in "ABCDEFGHI":
        ws[f"{col}{r}"].font = BODY
        ws[f"{col}{r}"].alignment = WRAP if col in "EFHI" else Alignment(vertical="top", horizontal="center" if col in "ABC" else None)
ws.column_dimensions["K"].hidden = True
ws.column_dimensions["J"].width = 2
ws.column_dimensions["G"].width = 22
ws.defined_names["_xlnm.Print_Area"] = DefinedName(
    "_xlnm.Print_Area", attr_text=f"OFFSET('Event Report'!$A$1,0,0,{R_FIRST - 1}+MAX(1,COUNT('Event Report'!$K${R_FIRST}:$K${R_LAST})),9)")
ws.print_title_rows = "17:17"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.oddFooter.left.text = "BOOST Ontario — initials only; contact details held separately"
ws.oddFooter.right.text = "Page &P of &N"
ws.conditional_formatting.add(f"A{R_FIRST}:I{R_LAST}", FormulaRule(formula=[f'AND($K{R_FIRST}<>"",MOD($A{R_FIRST},2)=0)'], fill=fill("F3F5F8"), border=BOX))
ws.conditional_formatting.add(f"A{R_FIRST}:I{R_LAST}", FormulaRule(formula=[f'$K{R_FIRST}<>""'], border=BOX))
protect(ws)

# =================================================================== Dashboard
ws = S["Dashboard"]
title(ws, "Dashboard", "Live totals. Nothing to type here.")
ws.column_dimensions["A"].width = 3
tiles = [("B", "People logged", f"=SUM({PZ})"),
         ("D", "Form sign-ups", f"=COUNT('New Sign-Ups'!$H${I_FIRST}:$H${I_LAST})"),
         ("F", "Not yet on People", "='New Sign-Ups'!$C$3"),
         ("H", "Events held", f'=COUNTIFS(EvDate,"<="&TODAY())'),
         ("J", "Follow-ups overdue", "='Follow-Ups'!$B$3"),
         ("L", "Follow-ups open", "='Follow-Ups'!$F$3")]
for col, lbl, frm in tiles:
    nxt = chr(ord(col) + 1)
    ws.merge_cells(f"{col}4:{nxt}4")
    ws.merge_cells(f"{col}5:{nxt}5")
    ws[f"{col}4"] = lbl
    ws[f"{col}4"].font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    ws[f"{col}4"].fill = fill(NAVY)
    ws[f"{col}4"].alignment = CENTER
    ws[f"{col}5"] = frm
    ws[f"{col}5"].font = Font(name="Calibri", size=22, bold=True, color="C00000" if "overdue" in lbl else NAVY)
    ws[f"{col}5"].fill = BANNER
    ws[f"{col}5"].alignment = CENTER
    ws.column_dimensions[col].width = 12
    ws.column_dimensions[nxt].width = 9
ws.row_dimensions[5].height = 38


def table(ws, top, left, heading, rows_, cols=("", "Count")):
    c0 = left
    c1 = chr(ord(left) + 1)
    ws[f"{c0}{top}"] = heading
    ws[f"{c0}{top}"].font = H2
    header(ws, top + 1, [(c0, cols[0] or heading.split(" ")[0], None, None), (c1, cols[1], None, None)], height=20)
    for i, (lbl, frm) in enumerate(rows_):
        r = top + 2 + i
        ws[f"{c0}{r}"] = lbl
        ws[f"{c1}{r}"] = frm
        for c in (c0, c1):
            ws[f"{c}{r}"].border = BOX
            ws[f"{c}{r}"].font = BODY
        ws[f"{c1}{r}"].alignment = Alignment(horizontal="center")
    return top + 2 + len(rows_)


pc = lambda col, val: f'=COUNTIFS(People!${col}${P_FIRST}:${col}${P_LAST},"{val}",{PZ},1)'  # noqa: E731
end1 = table(ws, 7, "B", "By contact type", [(v, pc("E", v)) for v in LISTS["L_Type"][1]], ("Contact Type", "People"))
end2 = table(ws, 7, "D", "By service level", [(v, pc("K", v)) for v in LISTS["L_Service"][1]], ("Service Level", "People"))
end3 = table(ws, 7, "F", "By stage", [(v, pc("M", v)) for v in LISTS["L_Stage"][1]], ("Stage", "People"))
end4 = table(ws, 7, "H", "By region", [(v, pc("G", v)) for v in LISTS["L_Region"][1]], ("Region", "People"))
for col, w in (("B", 32), ("D", 26), ("F", 16), ("H", 22)):
    ws.column_dimensions[col].width = w
for col in "CEGI":
    ws.column_dimensions[col].width = 10

top = max(end1, end2, end3, end4) + 2
ws[f"B{top}"] = "Progress against commitments"
ws[f"B{top}"].font = H2
ws[f"B{top+1}"] = ("People-tab counts come from rows you've logged. Self-identification counts come straight from the form "
                   "as totals — they are never copied onto anyone's row.")
ws[f"B{top+1}"].font = SUB
header(ws, top + 2, [("B", "Commitment", None, None), ("C", "Target", None, None), ("D", "So far", None, None),
                     ("E", "% of target", None, None), ("F", "Counted from", None, None)], height=20)
selfid = "INDEX(FormData,0,f_SelfID)"
kpis = [("KPI 1 — Total connections", f"=SUM({PZ})", "People logged"),
        ("KPI 2 — Underrepresented groups", f'=IFERROR(COUNTIF({selfid},"?*")-COUNTIF({selfid},"Prefer not to say"),"—")', "Form: self-ID answered (excl. prefer not to say)"),
        ("KPI 3 — Jobseekers", pc("H", "Jobseeker")[0:], "People → Situation"),
        ("KPI 4 — Incumbent workers", pc("H", "Incumbent worker"), "People → Situation"),
        ("KPI 6 — Registered apprentices", pc("I", "Yes"), "People → Reg. Apprentice"),
        ("KPI 26 — Persons with a disability", f'=IFERROR(COUNTIF({selfid},"*disabilit*"),"—")', "Form: self-ID"),
        ("Full-Service participants", pc("K", "Full-Service"), "People → Service Level"),
        ("Light-Touch participants", pc("K", "Light-Touch"), "People → Service Level")]
for i, (lbl, frm, src) in enumerate(kpis):
    r = top + 3 + i
    ws[f"B{r}"] = lbl
    ws[f"C{r}"] = f"=IF(Setup!$C${TARGET_ROW[lbl]}=\"\",\"—\",Setup!$C${TARGET_ROW[lbl]})"
    ws[f"D{r}"] = frm
    ws[f"E{r}"] = f'=IFERROR(D{r}/C{r},"")'
    ws[f"E{r}"].number_format = "0%"
    ws[f"F{r}"] = src
    for c in "BCDEF":
        ws[f"{c}{r}"].border = BOX
        ws[f"{c}{r}"].font = BODY if c != "D" else BOLD
    for c in "CDE":
        ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
    ws[f"F{r}"].font = Font(name="Calibri", size=9, italic=True, color="595959")
kp_first, kp_last = top + 3, top + 2 + len(kpis)
ws.conditional_formatting.add(f"E{kp_first}:E{kp_last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5B9BD5"))
ws.merge_cells(f"F{top+2}:I{top+2}")
for r in range(kp_first, kp_last + 1):
    ws.merge_cells(f"F{r}:I{r}")

top = kp_last + 3
ws[f"B{top}"] = "By event"
ws[f"B{top}"].font = H2
header(ws, top + 1, [("B", "Event", None, None), ("C", "Date", None, None), ("D", "Name", None, None),
                     ("E", "Sign-ups", None, None), ("F", "Logged", None, None), ("G", "Not yet logged", None, None),
                     ("H", "Follow-ups open", None, None)], height=30)
ev_first = top + 2
for i in range(60):
    r, er = ev_first + i, E_FIRST + i
    ws[f"B{r}"] = f'=IF(Events!$A{er}="","",Events!$A{er})'
    ws[f"C{r}"] = f'=IF(Events!$A{er}="","",IF(Events!$B{er}="","",Events!$B{er}))'
    ws[f"C{r}"].number_format = DATE
    ws[f"D{r}"] = f'=IF(Events!$A{er}="","",Events!$C{er}&"")'
    for c, src in zip("EFGH", "IJKL"):
        ws[f"{c}{r}"] = f'=IF(Events!$A{er}="","",Events!${src}{er})'
        ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
    for c in "BCDEFGH":
        ws[f"{c}{r}"].font = BODY
ev_last = ev_first + 59
ws.conditional_formatting.add(f"E{ev_first}:E{ev_last}", DataBarRule(start_type="num", start_value=0, end_type="max", color="5B9BD5"))
ws.conditional_formatting.add(f"F{ev_first}:F{ev_last}", DataBarRule(start_type="num", start_value=0, end_type="max", color="70AD47"))
ws.conditional_formatting.add(f"G{ev_first}:G{ev_last}", FormulaRule(formula=[f'N($G{ev_first})>0'], font=Font(color="C55A11", bold=True)))
ws.conditional_formatting.add(f"B{ev_first}:H{ev_last}", FormulaRule(formula=[f'$B{ev_first}<>""'], border=Border(bottom=THIN)))
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
protect(ws)

# =================================================================== Contact Card
ws = S["Contact Card"]
title(ws, "Contact Card", "The one place a person's name, email and phone appear. Type a Participant #, make the call, then clear the box.")
ws["A3"] = "🔒  Private — don't screen-share this tab or print it."
ws["A3"].font = Font(name="Calibri", size=11, bold=True, color="C00000")
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 44
ws["B5"] = "Participant # ▸"
ws["B5"].font = Font(name="Calibri", size=12, bold=True)
ws["B5"].alignment = Alignment(horizontal="right")
ws["C5"].fill = INPUT_FILL
ws["C5"].protection = UNLOCK
ws["C5"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
ws["C5"].border = Border(left=Side("medium", NAVY), right=Side("medium", NAVY), top=Side("medium", NAVY), bottom=Side("medium", NAVY))
ws["D5"] = '=IF(C5="","  e.g. P-0012",IF(F5="","  ✗ no such Participant #",IF(G5="","  ⚠ no form response linked — check People → Form Response ID","")))'
ws["D5"].font = Font(name="Calibri", size=10, italic=True, color="C00000")
ws["F5"] = f'=IF(C5="","",IFERROR(MATCH(UPPER(TRIM(C5)),People!$A${P_FIRST}:$A${P_LAST},0)+{P_FIRST - 1},""))'
ws["G5"] = '=IF(F5="","",INDEX(People!$W:$W,F5))'
ws.column_dimensions["F"].hidden = True
ws.column_dimensions["G"].hidden = True
card = [("Initials", '=IF($F$5="","",INDEX(People!$D:$D,$F$5)&"")'),
        ("Event", '=IF($F$5="","",INDEX(People!$B:$B,$F$5)&"")'),
        ("Consent to contact", '=IF($F$5="","",IF(INDEX(People!$T:$T,$F$5)="No","⚠ NO — do not contact",INDEX(People!$T:$T,$F$5)&""))'),
        ("Follow-up", '=IF($F$5="","",INDEX(People!$S:$S,$F$5)&IF(INDEX(People!$Q:$Q,$F$5)<>"","  ("&TEXT(INDEX(People!$Q:$Q,$F$5),"yyyy-mm-dd")&")",""))'),
        (None, None),
        ("Full name", '=IF($G$5="","",IFERROR(INDEX(FormData,$G$5,f_Name)&"",""))'),
        ("Email", '=IF($G$5="","",IFERROR(INDEX(FormData,$G$5,f_Email)&"",""))'),
        ("Phone", '=IF($G$5="","",IFERROR(INDEX(FormData,$G$5,f_Phone)&"",""))'),
        ("City", '=IF($G$5="","",IFERROR(INDEX(FormData,$G$5,f_City)&"",""))'),
        ("Organization", '=IF($F$5="","",INDEX(People!$F:$F,$F$5)&"")')]
for i, (lbl, frm) in enumerate(card):
    r = 7 + i
    if lbl is None:
        continue
    ws[f"B{r}"] = lbl
    ws[f"B{r}"].font = BOLD
    ws[f"B{r}"].alignment = Alignment(horizontal="right")
    ws[f"C{r}"] = frm
    ws[f"C{r}"].font = Font(name="Calibri", size=12)
    ws[f"C{r}"].border = Border(bottom=THIN)
ws.conditional_formatting.add("C9", FormulaRule(formula=['LEFT($C$9,1)="⚠"'], font=Font(color="C00000", bold=True)))
ws.print_area = "A1:A1"
protect(ws)

# =================================================================== Start Here
ws = S["Start Here"]
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 100
ws["B1"] = "BOOST Ontario — Participant Tracker"
ws["B1"].font = Font(name="Calibri", size=20, bold=True, color=NAVY)
ws["B2"] = "Literacy Link South Central  ·  Version 4  ·  October 2026"
ws["B2"].font = SUB
if DEMO:
    ws["B3"] = "SAMPLE COPY — every person in this file is made up. Use “BOOST Tracker v4.xlsx” for real work."
    ws["B3"].font = Font(name="Calibri", size=11, bold=True, color="C55A11")
    ws["B3"].fill = DEMO_FILL
    ws.merge_cells("B3:C3")
blocks = [
    ("THE IDEA", [
        ("Private by design", "People are tracked by Participant # and initials. Names, emails and phone numbers stay in the "
                              "Microsoft Forms response sheet and only appear on Contact Card, one person at a time."),
        ("One number links it", "Type the form's Response ID on a People row. Event, date, initials, type, region and consent fill themselves."),
        ("Notes are the point", "Contact details can always be looked up again. What someone told you at the table can't — that's what Conversation Notes is for."),
    ]),
    ("BEFORE THE EVENT", [
        ("1. Events", "Add a row: Event ID (EV-013, EV-014 …) and the DATE. Sign-ups find their event by the date, so the date matters most."),
    ]),
    ("AT THE TABLE", [
        ("2. They scan the QR", "The form does the capturing. Nothing to type in the tracker yet."),
        ("Walk-ups", "No form? Add a People row: type the date in Registered On, their initials, and the Contact Type. The Event ID fills from the date."),
    ]),
    ("AFTER THE EVENT", [
        ("3. New Sign-Ups", "Shows everyone who filled in the form. Anyone marked → add isn't on People yet."),
        ("4. People", "For each one, type the Response ID in column U on the next empty row. Then add Service Level, Main Need, Stage and your notes."),
        ("5. Follow-Ups", "Tick Follow-Up Needed and give it a due date. Follow-Ups lists everyone you owe a call, soonest first. Enter the date in Follow-Up Done when it's done."),
        ("6. Event Report", "Pick the Event ID, write your note to the reader, then File → Print → Save as PDF. It shows initials only, so it's safe to send."),
    ]),
    ("WHEN YOU NEED TO CALL", [
        ("Contact Card", "Type the Participant #. Name, email and phone appear. Clear the box afterwards. Never screen-share this tab."),
    ]),
    ("COLOURS", [
        ("Header colours", "Teal = fills itself from the form, but you can type over it.  Navy = yours.  Grey = counts itself.  Orange = the link (Form Response ID)."),
        ("Cell colours", "Yellow = type here.  Red on Event ID = no event matched that date.  Amber on Event ID = two events on the same day, so check it.  "
                         "Light grey = doesn't apply to this contact type."),
    ]),
    ("PRIVACY RULES", [
        ("In notes", "Write what they need and what you agreed, not who they are. No health details, diagnoses, SIN, dates of birth or addresses."),
        ("Self-identification", "Never typed into the tracker. The Dashboard counts it straight from the form as totals only."),
        ("Consent = No", "Shows in red on People and Contact Card. Don't contact them."),
        ("Sharing", "People, Events, Follow-Ups, Event Report and Dashboard hold no contact details. Contact Card and Sheet1 do."),
    ]),
    ("CONNECTING THE FORM", [
        ("Where responses live", "Microsoft Forms writes responses to “Event Registration - Literacy Link South Central.xlsx” (OneDrive → Documents), on a sheet called Sheet1."),
        ("Option A — paste (simplest)", "In Forms: Responses → Open in Excel. Select all on that sheet, copy, and paste over this file's Sheet1 at A1. Takes 30 seconds after each event."),
        ("Option B — live", "Copy every tab EXCEPT Sheet1 into the Forms workbook (right-click a tab → Move or Copy → tick Create a copy; do Lists and Setup first). "
                            "Then Data → Edit Links → Change Source → choose the Forms workbook itself, so the formulas point at its real Sheet1."),
        ("Setup tab", "If a heading on the form is worded differently, Setup will show ✗. Type a word from the real heading into the yellow cell and it reconnects."),
    ]),
    ("DON'T", [
        ("…", "Don't delete or rename Sheet1.  Don't delete rows on People — clear them instead (Participant #s must stay unique).  "
              "Don't unhide and edit the grey helper columns (W–Z on People)."),
    ]),
]
r = 5
for head, items in blocks:
    ws[f"B{r}"] = head
    ws[f"B{r}"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws[f"B{r}"].fill = fill(NAVY)
    ws[f"C{r}"].fill = fill(NAVY)
    r += 1
    for lbl, txt in items:
        ws[f"B{r}"] = lbl
        ws[f"B{r}"].font = BOLD
        ws[f"B{r}"].alignment = WRAP
        ws[f"C{r}"] = txt
        ws[f"C{r}"].font = BODY
        ws[f"C{r}"].alignment = WRAP
        ws.row_dimensions[r].height = 15 * max(1, -(-len(txt) // 105))
        r += 1
    r += 1
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# =================================================================== tab colours
for n, c in {"Start Here": NAVY, "People": NAVY, "New Sign-Ups": "C55A11", "Follow-Ups": "C00000",
             "Events": TEAL, "Event Report": TEAL, "Dashboard": TEAL, "Contact Card": "7F7F7F",
             "Setup": "BFBFBF", "Lists": "BFBFBF", "Sheet1": "BFBFBF"}.items():
    S[n].sheet_properties.tabColor = c

# =================================================================== Events known so far (both files)
KNOWN_EVENTS = [("EV-001", "SOY Boot Camp — Ottawa", "SOY Boot Camp", "Ottawa", "Eastern Ontario"),
                ("EV-002", "SOY Boot Camp — Markham", "SOY Boot Camp", "Markham", "Greater Toronto Area"),
                ("EV-004", "SOY Boot Camp — Brampton", "SOY Boot Camp", "Brampton", "Greater Toronto Area"),
                ("EV-006", "SOY Boot Camp — Alliston", "SOY Boot Camp", "Alliston", "Central Ontario"),
                ("EV-008", "SOY Boot Camp — Cobourg", "SOY Boot Camp", "Cobourg", "Eastern Ontario"),
                ("EV-009", "Level Up — Cobourg", "Level Up", "Cobourg", "Eastern Ontario"),
                ("EV-010", "SOY Boot Camp — Oshawa", "SOY Boot Camp", "Oshawa", "Greater Toronto Area"),
                ("EV-011", "Level Up — Kingston", "Level Up", "Kingston", "Eastern Ontario"),
                ("EV-012", "SOY Boot Camp — Acton", "SOY Boot Camp", "Acton", "Central Ontario")]

# =================================================================== demo data
if DEMO:
    random.seed(20261009)
    ev_ws, ppl, form = S["Events"], S["People"], S["Sheet1"]
    demo_dates = [dt.date(2026, 9, 9), dt.date(2026, 9, 11), dt.date(2026, 9, 16), dt.date(2026, 9, 18),
                  dt.date(2026, 9, 23), dt.date(2026, 9, 24), dt.date(2026, 9, 30), dt.date(2026, 10, 2),
                  dt.date(2026, 10, 7)]
    hosts = ["Skills Ontario", "Skills Ontario", "Skills Ontario", "Skills Ontario", "Skills Ontario",
             "Northumberland CFDC", "Skills Ontario", "St. Lawrence College", "Skills Ontario"]
    worked = ["Busy table at lunch. Postcards ran out by 1 pm, so bring 150 next time.",
              "Students were keen. Teachers asked for a classroom visit.",
              "Quiet morning, strong afternoon. The QR stand at eye level helped.",
              "Small group, but three employers stayed to talk sponsorship.",
              "Partnered with the local employment centre. Good referrals both ways.",
              "Evening session. Mostly incumbent workers asking about exam prep.",
              "Big turnout. Need a second person at the table.",
              "College hosted. Mostly 1st-year apprentices.",
              ""]
    for i, ((eid, nm, typ, city, reg), d) in enumerate(zip(KNOWN_EVENTS, demo_dates)):
        r = E_FIRST + i
        ev_ws[f"A{r}"], ev_ws[f"B{r}"], ev_ws[f"C{r}"], ev_ws[f"D{r}"] = eid, d, nm, typ
        ev_ws[f"E{r}"], ev_ws[f"F{r}"], ev_ws[f"G{r}"], ev_ws[f"H{r}"] = hosts[i], city, reg, "M. Caldwell"
        ev_ws[f"N{r}"] = worked[i]
        if d < dt.date(2026, 10, 5):
            ev_ws[f"M{r}"] = d + dt.timedelta(days=2)
    r = E_FIRST + len(KNOWN_EVENTS)
    ev_ws[f"A{r}"], ev_ws[f"B{r}"], ev_ws[f"C{r}"], ev_ws[f"D{r}"] = "EV-013", dt.date(2026, 10, 21), "Career Fair — Peterborough", "Career fair"
    ev_ws[f"E{r}"], ev_ws[f"F{r}"], ev_ws[f"G{r}"], ev_ws[f"H{r}"] = "Fleming College", "Peterborough", "Central Ontario", "M. Caldwell"

    firsts = ["Alex", "Jordan", "Priya", "Sam", "Taylor", "Morgan", "Chris", "Dana", "Jamie", "Riley", "Avery", "Casey",
              "Noah", "Liam", "Emma", "Olivia", "Ethan", "Maya", "Lucas", "Zoe", "Omar", "Fatima", "Wei", "Mateo",
              "Hannah", "Isaac", "Nadia", "Tariq", "Grace", "Owen", "Leah", "Kai", "Sofia", "Ben", "Aisha", "Marc"]
    lasts = ["Martin", "Tremblay", "Singh", "Roy", "Brown", "Wilson", "Gagnon", "Lee", "Patel", "Campbell", "Nguyen",
             "Bouchard", "Thompson", "Khan", "Clarke", "Moreau", "Ali", "Fraser", "Wong", "Côté", "Hughes", "Bell"]
    who_ans = {"Trade": "I'm in a trade, or thinking about one", "Employer & Union": "I'm an employer or union rep",
               "College & Agency": "I work for a college, agency or community organization", "Other": "Other / just curious"}
    JS, IW, JP, ST = ("I'm not working right now and I'm looking for work", "I'm working, but I'm not a journeyperson",
                      "I'm a journeyperson", "I'm a student")
    # situation, registered apprentice, trade, form "most helpful", Main Need, Service Level, note
    TRADE = [
        (JS, "Yes", "Electrician", "Finding a sponsor or employer", "Finding a sponsor / employer", "Full-Service",
         "Finished level 1 in 2024, then lost the sponsor when the shop closed. Wants a new sponsor close to home."),
        (IW, "Yes", "Plumber", "Help preparing for my exam", "Exam preparation", "Full-Service",
         "Writing the C of Q exam in November. Failed once on code questions. Interested in tutoring twice a week."),
        (ST, "No", "Not sure yet", "Just want to know more about BOOST", "Info about BOOST", "Light-Touch",
         "Grade 11. Parent was with them. Curious about electrical vs plumbing. Took a postcard home."),
        (IW, "Not sure", "Carpenter", "Figuring out where to start", "Where to start", "Full-Service",
         "Has worked as a helper for 2 years. Not sure the training agreement was ever registered. Walked them through the steps."),
        (JS, "No", "Welder", "Help preparing for my exam", "Exam preparation", "Full-Service",
         "Newcomer, licensed welder back home. Asked about trade equivalency and the exam. Pointed to Skills Ontario's TEA info."),
        (JS, "No", "Not sure yet", "Tutoring or upgrading", "Tutoring / upgrading", "Light-Touch",
         "Wants a math upgrade before applying to pre-apprenticeship. Evenings only, works days."),
        (IW, "Yes", "Automotive Service Technician", "Help with reading, writing or math", "Learning disability testing", "Full-Service",
         "Says reading the exam questions is the hard part. Asked about a learning disability assessment."),
        (JS, "No", "HVAC", "Finding a sponsor or employer", "Finding a sponsor / employer", "Full-Service",
         "Very keen. Has an employer willing to sponsor but needs help with the paperwork."),
        (ST, "No", "Cook", "Just want to know more about BOOST", "Info about BOOST", "Light-Touch",
         "Just browsing, took a card. Low interest for now."),
        (IW, "Yes", "Millwright", "Tutoring or upgrading", "Tutoring / upgrading", "Full-Service",
         "3rd-year apprentice struggling with the in-school portion. Wants to talk through tutoring options."),
        (JP, "No", "Electrician", "Just want to know more about BOOST", "Info about BOOST", "Light-Touch",
         "Licensed electrician. Interested in mentoring apprentices and asked how to get involved."),
        (JS, "No", "Hairstylist", "Figuring out where to start", "Where to start", "Light-Touch",
         "Did the in-school part years ago but never finished hours. Wants to know if they can pick it back up."),
    ]
    EMP = [("Northshore Electric Ltd.", "Hiring an apprentice",
            "Owner of a 12-person shop. Would sponsor one apprentice in spring if the wage subsidy is confirmed. Send the info sheet."),
           ("IBEW training centre (sample)", "Partnership / referrals",
            "Union training coordinator. Wants to refer members who've failed the exam. Asked for a short presentation."),
           ("Kawartha Mechanical", "Hiring an apprentice",
            "HR lead. Two incumbent workers close to the exam. Asked if we can run a group session on site.")]
    ORG = [("Durham District School Board", "Guidance counsellor. Wants a class visit in November: grades 11–12, about 60 students."),
           ("YMCA Employment Services", "Employment counsellor. Agreed on two-way referrals. Will send clients who need exam prep.")]
    OTHER_NOTE = "Parent asking for their son (17). Took a postcard."
    selfids = ["", "", "", "Youth (15–29)", "Woman", "Woman;Youth (15–29)", "Newcomer or immigrant",
               "Person with a disability", "Black or racialized person;Youth (15–29)", "Francophone", "Prefer not to say",
               "Indigenous (First Nations, Inuit, Métis)", "Person with a disability;Woman"]

    resp_id = 0
    responses = []
    per_event = [7, 5, 9, 4, 6, 5, 8, 6, 3]
    counters = {"Trade": 0, "Employer & Union": 0, "College & Agency": 0, "Other": 0}
    for (eid, _, _, _, reg), d, n in zip(KNOWN_EVENTS, demo_dates, per_event):
        for _ in range(n):
            resp_id += 1
            t = random.random()
            typ = "Trade" if t < 0.68 else "Employer & Union" if t < 0.82 else "College & Agency" if t < 0.93 else "Other"
            k = counters[typ]
            counters[typ] += 1
            when = dt.datetime.combine(d, dt.time(10, 0)) + dt.timedelta(minutes=random.randint(0, 330))
            responses.append(dict(id=resp_id, event=eid, date=d, when=when, typ=typ, k=k, region=reg,
                                  fn=random.choice(firsts), ln=random.choice(lasts)))
    # one late response: took a card at Kingston, filled in the form at home three days later
    resp_id += 1
    responses.append(dict(id=resp_id, event=None, date=dt.date(2026, 10, 5), when=dt.datetime(2026, 10, 5, 20, 14),
                          typ="Trade", k=1, region="Eastern Ontario", fn="Riley", ln="Fraser"))
    for i, rsp in enumerate(responses):
        typ, k = rsp["typ"], rsp["k"]
        p = TRADE[k % len(TRADE)] if typ == "Trade" else None
        vals = [rsp["id"], rsp["when"] - dt.timedelta(minutes=2), rsp["when"], "anonymous", "", rsp["when"],
                who_ans[typ], f"{rsp['fn']} {rsp['ln']}", f"{rsp['fn'].lower()}.{rsp['ln'].lower()}@example.com",
                f"555-01{rsp['id']:02d}", "Sample City",
                EMP[k % len(EMP)][0] if typ == "Employer & Union" else "",
                ORG[k % len(ORG)][0] if typ == "College & Agency" else "",
                rsp["region"],
                p[0] if p else "", p[1] if p else "", p[2] if p else "",
                random.choice(selfids) if p else "", p[3] if p else "",
                "No" if rsp["id"] in (19, 44) else
                "Yes, I agree to be contacted by Literacy Link South Central about BOOST Ontario"]
        for j, v in enumerate(vals, 1):
            c = form.cell(row=i + 2, column=j, value=v)
            if isinstance(v, dt.datetime):
                c.number_format = DATETIME

    logged = [r for r in responses if not (r["event"] == "EV-012" or r["id"] in (14, 31))]  # leave a few unlogged
    pr = P_FIRST
    for rsp in logged:
        typ, k = rsp["typ"], rsp["k"]
        ppl[f"U{pr}"] = rsp["id"]
        if rsp["event"] is None:
            ppl[f"B{pr}"] = "EV-011"  # typed over: the date matched no event
            ppl[f"N{pr}"] = "Took a card at Kingston and filled in the form at home. "
        note = ppl[f"N{pr}"].value or ""
        if typ == "Trade":
            p = TRADE[k % len(TRADE)]
            ppl[f"K{pr}"], ppl[f"L{pr}"], ppl[f"N{pr}"] = p[5], p[4], note + p[6]
        elif typ == "Employer & Union":
            ppl[f"L{pr}"], ppl[f"N{pr}"] = EMP[k % len(EMP)][1], EMP[k % len(EMP)][2]
        elif typ == "College & Agency":
            ppl[f"L{pr}"], ppl[f"N{pr}"] = "Partnership / referrals", ORG[k % len(ORG)][1]
        else:
            ppl[f"L{pr}"], ppl[f"N{pr}"] = "Info about BOOST", OTHER_NOTE
        if random.random() < 0.7:
            ppl[f"O{pr}"] = "✓"
        low_interest = typ == "Other" or (typ == "Trade" and TRADE[k % len(TRADE)][5] == "Light-Touch" and random.random() < 0.5)
        if rsp["id"] in (19, 44):
            ppl[f"M{pr}"] = "Closed"
            ppl[f"N{pr}"] = "Said no to contact on the form. Information only. " + ppl[f"N{pr}"].value
        elif low_interest:
            ppl[f"M{pr}"] = "Closed" if (TODAY - rsp["date"]).days > 14 else "New"
        else:
            ppl[f"P{pr}"] = "✓"
            due = rsp["date"] + dt.timedelta(days=random.choice([5, 7, 10, 14, 21]))
            ppl[f"Q{pr}"] = due
            if due < TODAY and random.random() < 0.7:
                ppl[f"R{pr}"] = due - dt.timedelta(days=random.randint(0, 3))
                ppl[f"M{pr}"] = random.choice(["Contacted", "Intake booked", "Active", "Active", "Referred out"])
            elif due < TODAY:
                ppl[f"M{pr}"] = random.choice(["New", "No response"])
            else:
                ppl[f"M{pr}"] = "New" if (TODAY - rsp["date"]).days < 7 else "Contacted"
        pr += 1
    # two walk-ups with no form response
    for when, ini, typ, need, note in [
            (dt.date(2026, 9, 16), "DR", "Employer & Union", "Hiring an apprentice",
             "Walk-up, didn't want to use the QR. Runs a small HVAC company. Took a card and will email us."),
            (dt.date(2026, 9, 30), "KM", "Trade", "Exam preparation",
             "Walk-up whose phone was dead, so no form. Writing the exam in January. Asked them to fill in the form from home.")]:
        ppl[f"C{pr}"], ppl[f"D{pr}"], ppl[f"E{pr}"], ppl[f"L{pr}"], ppl[f"N{pr}"] = when, ini, typ, need, note
        ppl[f"M{pr}"] = "Contacted"
        ppl[f"P{pr}"], ppl[f"Q{pr}"] = "✓", TODAY + dt.timedelta(days=3)
        pr += 1
    S["Event Report"]["C3"] = "EV-010"
    S["Event Report"]["B15"] = ("Good turnout: eight sign-ups plus one walk-up. Most were working in a trade or looking for a way back in, "
                                "and several asked about exam prep. One employer wants an on-site group session for two staff. "
                                "Follow-ups are booked over the next two weeks.")
    S["Contact Card"]["C5"] = None
else:
    for i, (eid, nm, typ, city, reg) in enumerate(KNOWN_EVENTS):
        r = E_FIRST + i
        e = S["Events"]
        e[f"A{r}"], e[f"C{r}"], e[f"D{r}"], e[f"F{r}"], e[f"G{r}"] = eid, nm, typ, city, reg
    S["Events"]["N3"] = ("Filled in from your Event Report filenames. Add each DATE (red until you do). "
                         "EV-003, EV-005 and EV-007 had no report, so add them if they happened.")
    S["Events"]["N3"].font = Font(name="Calibri", size=9, italic=True, color="C55A11")
    S["Events"]["N3"].alignment = WRAP
    S["Events"].row_dimensions[3].height = 36

# =================================================================== document properties
wb.properties.creator = "Maxime Caldwell"
wb.properties.lastModifiedBy = "Maxime Caldwell"
wb.properties.title = "BOOST Ontario — Participant Tracker"
wb.properties.subject = "BOOST Ontario"
wb.properties.company = None
wb.properties.created = wb.properties.modified = dt.datetime(2026, 10, 8, 22, 0)
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)

# plain application stamp, as Excel itself writes it
import shutil, zipfile  # noqa: E401,E402
tmp = OUT.with_suffix(".tmp")
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "docProps/app.xml":
            data = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
                    '<Application>Microsoft Excel</Application><DocSecurity>0</DocSecurity>'
                    '<AppVersion>16.0300</AppVersion></Properties>').encode()
        zout.writestr(item, data)
shutil.move(tmp, OUT)
print("wrote", OUT)
